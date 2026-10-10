"""File parsers and skill extraction heuristics for candidate documents."""
import hashlib
import json
from pathlib import Path
import re
from typing import Dict, List, Tuple, Any, Optional

import pypdf
import docx


COMMON_SKILLS_CATALOG = {
    "data_engineering": [
        "pyspark", "spark", "databricks", "azure data factory", "adf", "azure synapse",
        "dbt", "airflow", "delta lake", "snowflake", "sql", "transact-sql", "t-sql",
        "kafka", "event hubs", "etl", "elt", "data lake", "data warehouse", "lakehouse",
        "pandas", "polars", "glue", "redshift", "bigquery", "data modeling", "star schema"
    ],
    "cloud_infrastructure": [
        "azure", "aws", "gcp", "docker", "kubernetes", "k8s", "terraform", "bicep",
        "arm templates", "ci/cd", "azure devops", "github actions", "linux", "bash",
        "powershell", "helm"
    ],
    "programming": [
        "python", "scala", "java", "sql", "c#", ".net", "javascript", "typescript",
        "golang", "rust", "r"
    ],
    "databases": [
        "postgresql", "mysql", "sql server", "ssms", "mongodb", "redis", "cosmos db",
        "dynamodb", "cassandra"
    ],
    "data_governance_bi": [
        "power bi", "powerbi", "tableau", "looker", "purview", "data quality", "great expectations"
    ],
    "architecture_methodologies": [
        "system design", "microservices", "clean architecture", "scrum", "agile",
        "kanban", "devops", "git", "medallion architecture"
    ]
}


def compute_file_hash(path: Path) -> str:
    """Computes SHA-256 hash of a file."""
    hasher = hashlib.sha256()
    with open(path, "rb") as f:
        while chunk := f.read(65536):
            hasher.update(chunk)
    return hasher.hexdigest()


def extract_text_from_pdf(path: Path) -> str:
    """Extracts text from a PDF document safely."""
    text_parts = []
    try:
        reader = pypdf.PdfReader(str(path))
        for page_num, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            if page_text.strip():
                text_parts.append(page_text.strip())
    except Exception as e:
        return f"[Error extracting PDF text from {path.name}: {e}]"
    return "\n\n".join(text_parts)


def extract_text_from_docx(path: Path) -> str:
    """Extracts text from a Word (.docx) document safely."""
    text_parts = []
    try:
        doc = docx.Document(str(path))
        for p in doc.paragraphs:
            if p.text.strip():
                text_parts.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)
    except Exception as e:
        return f"[Error extracting DOCX text from {path.name}: {e}]"
    return "\n".join(text_parts)


def extract_text_from_txt(path: Path) -> str:
    """Extracts text from text, markdown, or json files."""
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()
    except Exception as e:
        return f"[Error extracting text from {path.name}: {e}]"


def parse_document_file(path: Path) -> Tuple[str, Dict[str, Any]]:
    """Extracts plain text and basic metadata according to file extension."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        raw_text = extract_text_from_pdf(path)
    elif suffix in (".docx", ".doc"):
        raw_text = extract_text_from_docx(path)
    elif suffix in (".txt", ".md", ".json", ".markdown"):
        raw_text = extract_text_from_txt(path)
    else:
        raw_text = extract_text_from_txt(path)

    metadata = {
        "file_name": path.name,
        "extension": suffix,
        "char_count": len(raw_text),
        "word_count": len(raw_text.split())
    }
    return raw_text, metadata


def extract_skills_from_text(text: str) -> List[Dict[str, str]]:
    """Extracts matching skills from text using catalog matching."""
    text_lower = text.lower()
    found_skills: List[Dict[str, str]] = []
    seen = set()

    for category, skills in COMMON_SKILLS_CATALOG.items():
        for skill in skills:
            # Word boundary regex search
            pattern = r"(?<!\w)" + re.escape(skill) + r"(?!\w)"
            if re.search(pattern, text_lower):
                if skill not in seen:
                    seen.add(skill)
                    found_skills.append({
                        "name": skill.title(),
                        "category": category,
                        "proficiency": "advanced" if text_lower.count(skill) > 3 else "intermediate"
                    })

    return found_skills
